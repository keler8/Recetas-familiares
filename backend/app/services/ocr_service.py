import os
import logging
from typing import Optional, List
from PIL import Image
import pytesseract
import pdfplumber
import docx
from app.services.storage_guard import storage_guard

logger = logging.getLogger(__name__)


class OCRService:
    @staticmethod
    def extract_text_from_image(image_path: str, lang: str = "spa") -> str:
        """Extracts text from an image using local Tesseract OCR."""
        try:
            with Image.open(image_path) as img:
                text = pytesseract.image_to_string(img, lang=lang)
                return text.strip()
        except Exception as e:
            logger.error(f"Error in OCR image processing {image_path}: {e}")
            return ""

    @staticmethod
    def extract_text_from_pdf(pdf_path: str) -> str:
        """
        Extracts digital text using pdfplumber, falling back to page OCR if pages are images.
        """
        extracted_pages = []
        try:
            with pdfplumber.open(pdf_path) as pdf:
                for page_idx, page in enumerate(pdf.pages):
                    text = page.extract_text()
                    if text and len(text.strip()) > 30:
                        extracted_pages.append(text.strip())
                    else:
                        # Page is likely a scanned image, attempt OCR
                        try:
                            page_image = page.to_image(resolution=200).original
                            ocr_text = pytesseract.image_to_string(page_image, lang="spa")
                            if ocr_text.strip():
                                extracted_pages.append(ocr_text.strip())
                        except Exception as ocr_err:
                            logger.warning(f"Could not OCR page {page_idx} of {pdf_path}: {ocr_err}")
            return "\n\n--- PÁGINA ---\n\n".join(extracted_pages)
        except Exception as e:
            logger.error(f"Error extracting text from PDF {pdf_path}: {e}")
            return ""

    @staticmethod
    def extract_text_from_docx(docx_path: str) -> str:
        """Extracts paragraphs and tables from a Word DOCX document."""
        try:
            doc = docx.Document(docx_path)
            full_text = []
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        full_text.append(row_text)
            return "\n".join(full_text)
        except Exception as e:
            logger.error(f"Error reading DOCX {docx_path}: {e}")
            return ""

    @classmethod
    def process_document(cls, file_path: str, mime_type: Optional[str] = None) -> str:
        """Master dispatcher for document extraction with automatic safety checks."""
        if not storage_guard.can_start_heavy_task():
            raise RuntimeError("Espacio de disco insuficiente para procesar documentos grandes.")

        ext = os.path.splitext(file_path)[1].lower()
        if ext in [".pdf"]:
            return cls.extract_text_from_pdf(file_path)
        elif ext in [".docx"]:
            return cls.extract_text_from_docx(file_path)
        elif ext in [".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tiff"]:
            return cls.extract_text_from_image(file_path)
        elif ext in [".txt", ".md"]:
            try:
                with open(file_path, "r", encoding="utf-8") as f:
                    return f.read()
            except Exception:
                with open(file_path, "r", encoding="latin-1") as f:
                    return f.read()
        else:
            raise ValueError(f"Formato de archivo no soportado para extracción: {ext}")


ocr_service = OCRService()
